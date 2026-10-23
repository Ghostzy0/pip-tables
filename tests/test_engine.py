import csv
import datetime as dt
import json
import threading
import time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from pip_tables.engine import Project


@pytest.fixture
def project(tmp_path):
    item = Project.create(tmp_path / "workspace")
    yield item
    item.close()


def test_csv_numeric_profile_join_history_reopen(project, tmp_path):
    sales = tmp_path / "sales.csv"
    sales.write_text('sku,amount,day,note\nA,10.25,2026-01-01,"a, b"\nB,2.50,2026-01-02,"line\nbreak"\nA,,2026-01-03,éclair\n', encoding="utf-8")
    targets = tmp_path / "targets.csv"
    targets.write_text("sku,target\nA,12\nB,5\n", encoding="utf-8")
    assert project.import_file(sales) == "sales"
    assert project.import_file(targets) == "targets"
    assert project.import_file(sales) == "sales_2"
    result = project.query("SELECT s.sku, sum(s.amount) AS actual, max(t.target) AS target "
                           "FROM sales s JOIN targets t USING (sku) GROUP BY s.sku ORDER BY s.sku")
    assert result.columns == ["sku", "actual", "target"]
    assert result.rows == [("A", pytest.approx(10.25), 12), ("B", pytest.approx(2.5), 5)]
    assert result.total == 2 and not result.truncated
    assert project.query("SELECT note FROM sales ORDER BY sku, day", limit=1).truncated
    profile = {item["name"]: item for item in project.profile("sales")}
    assert profile["amount"]["nulls"] == 1
    assert profile["sku"]["distinct"] == 2
    assert project.tables()[0]["source"] == "sales.csv"
    assert project.tables()[0]["rows"] == 3
    project.save_draft("SELECT * FROM sales")
    path = project.path
    project.close()
    reopened = Project(path)
    assert reopened.draft == "SELECT * FROM sales"
    assert len(reopened.history()) == 2
    assert reopened.query("SELECT count(*) FROM sales").rows == [(3,)]
    reopened.close()


def test_readonly_sql_and_failed_import_leave_project_intact(project, tmp_path):
    good = tmp_path / "good.csv"
    good.write_text("value\n7\n")
    project.import_file(good)
    before = project.tables()
    for sql in (
        "DELETE FROM good", "DROP TABLE good", "CREATE TABLE bad AS SELECT 1",
        "UPDATE good SET value=0", "COPY good TO '/tmp/pip-secret.csv'",
        "SELECT * FROM good; DELETE FROM good", "PRAGMA enable_external_access=true",
        "INSTALL httpfs", "LOAD httpfs", "SELECT * FROM read_csv('/etc/passwd')",
        "SELECT * FROM read_parquet('https://example.com/data.parquet')",
    ):
        with pytest.raises((ValueError, duckdb_error())):
            project.query(sql)
    assert project.tables() == before
    assert project.history() == []
    bad = tmp_path / "bad.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["value"])
    sheet.append([1])
    sheet.append(["=1+1"])
    book.save(bad)
    with pytest.raises(ValueError, match="formula"):
        project.import_file(bad)
    assert project.tables() == before


def duckdb_error():
    # Engine may expose DuckDB's specific permission error for a blocked
    # table function; all statement mutations must remain impossible.
    import duckdb
    return duckdb.Error


def test_xlsx_sheets_types_and_formula_rejection(project, tmp_path):
    path = tmp_path / "ledger.xlsx"
    book = Workbook()
    one = book.active
    one.title = "First"
    one.append(["same", "same", "date", "flag"])
    one.append([1, 2.5, dt.date(2026, 1, 2), True])
    one.append([3, None, None, False])
    second = book.create_sheet("Other")
    second.append(["value"])
    second.append([42])
    book.save(path)
    with pytest.raises(ValueError, match="Select an XLSX sheet"):
        project.import_file(path)
    name = project.import_file(path, sheet="First")
    assert [c["name"] for c in project.tables()[0]["columns"]] == ["same", "same_2", "date", "flag"]
    assert project.query(f'SELECT sum(same), sum("same_2") FROM "{name}"').rows == [(4, 2.5)]
    assert project.import_file(path, sheet="Other", all_text=True) == "ledger_2"
    assert project.query("SELECT value FROM ledger_2").rows == [("42",)]
    with pytest.raises(ValueError, match="sheet not found"):
        project.import_file(path, sheet="Absent")


def test_json_parquet_and_roundtrip_exports(project, tmp_path):
    data = tmp_path / "events.jsonl"
    data.write_text('{"id":1,"label":"alpha"}\n{"id":2,"label":"βeta"}\n', encoding="utf-8")
    name = project.import_file(data)
    assert project.query(f"SELECT sum(id) FROM {name}").rows == [(3,)]
    sql = "SELECT id, label FROM events ORDER BY id"
    for kind in ("csv", "xlsx", "json", "parquet"):
        destination = tmp_path / f"out_{kind}.{kind}"
        assert project.export(sql, destination, kind) == 2
        assert destination.is_file() and destination.stat().st_size > 0
        imported = project.import_file(destination)
        assert project.query(f"SELECT id, label FROM {imported} ORDER BY id").rows == [(1, "alpha"), (2, "βeta")]
    assert json.loads((tmp_path / "out_json.json").read_text()) == [
        {"id": 1, "label": "alpha"}, {"id": 2, "label": "βeta"}]


def test_formula_text_is_inert_in_csv_and_xlsx(project, tmp_path):
    source = tmp_path / "attacks.csv"
    source.write_text('payload\n"=1+1"\n"+SUM(1,2)"\n"@cmd"\n" -2"\n', encoding="utf-8")
    project.import_file(source, all_text=True)
    sql = "SELECT payload FROM attacks"
    csv_file = tmp_path / "safe.csv"
    xlsx_file = tmp_path / "safe.xlsx"
    assert project.export(sql, csv_file, "csv") == 4
    assert project.export(sql, xlsx_file, "xlsx") == 4
    with csv_file.open(newline="", encoding="utf-8") as stream:
        values = [row[0] for row in list(csv.reader(stream))[1:]]
    assert values == ["'=1+1", "'+SUM(1,2)", "'@cmd", "' -2"]
    book = load_workbook(xlsx_file, read_only=True, data_only=False)
    try:
        assert [row[0].value for row in list(book.active.iter_rows())[1:]] == values
        assert all(row[0].data_type != "f" for row in list(book.active.iter_rows())[1:])
    finally:
        book.close()


def test_duplicate_result_labels_survive_json_and_parquet(project, tmp_path):
    sql = "SELECT 1 AS label, 2 AS label"
    expected = {"label": 1, "label_1": 2}
    target = tmp_path / "duplicates.json"
    project.export(sql, target, "json")
    assert json.loads(target.read_text()) == [expected]
    parquet = tmp_path / "duplicates.parquet"
    project.export(sql, parquet, "parquet")
    name = project.import_file(parquet)
    assert project.query(f"SELECT label, label_1 FROM {name}").rows == [(1, 2)]


def test_invalid_project_and_limits(tmp_path, project):
    missing = tmp_path / "missing"
    with pytest.raises(FileNotFoundError):
        Project(missing)
    with pytest.raises(FileExistsError):
        Project.create(project.path)
    with pytest.raises(ValueError):
        project.query("SELECT 1", limit=0)
    with pytest.raises(ValueError):
        project.query("SELECT 1", timeout=0)
    (tmp_path / "unknown.txt").write_text("unsupported")
    with pytest.raises(ValueError):
        project.import_file(tmp_path / "unknown.txt")
    corrupt = tmp_path / "corrupt"
    corrupt.mkdir()
    (corrupt / "project.json").write_text("{")
    (corrupt / "data.duckdb").write_bytes(b"not a database")
    original = (corrupt / "data.duckdb").read_bytes()
    with pytest.raises(ValueError):
        Project(corrupt)
    assert (corrupt / "data.duckdb").read_bytes() == original
    assert project.query("SELECT 1").rows == [(1,)]
    project.close()
    project.close()
    with pytest.raises(RuntimeError):
        project.tables()


def test_project_rejects_unregistered_view(tmp_path):
    import duckdb
    project = Project.create(tmp_path / "p")
    database = project.path / "data.duckdb"
    project.close()
    connection = duckdb.connect(str(database))
    connection.execute("CREATE VIEW outside AS SELECT * FROM read_csv('/etc/passwd')")
    connection.close()
    with pytest.raises(ValueError, match="unrecognised"):
        Project(tmp_path / "p")


def test_bounded_preview_over_large_relation(project):
    result = project.query("SELECT i, i * i AS square FROM range(100000) AS t(i) ORDER BY i", limit=5)
    assert result.rows == [(i, i * i) for i in range(5)]
    assert result.total == 100_000 and result.truncated
    assert len(result.rows) == 5
    pivot = project.query("PIVOT (SELECT 'A' AS category, 7 AS amount) ON category USING sum(amount)")
    assert pivot.columns == ["A"] and pivot.rows == [(7,)]


def test_preview_rejects_large_cells_and_result_bytes_before_python_copy(project):
    with pytest.raises(ValueError, match="Preview exceeds"):
        project.query("SELECT repeat('x', 2000000) AS value")
    with pytest.raises(ValueError, match="Preview exceeds"):
        project.query("SELECT repeat('x', 100000) AS value FROM range(100)")
    assert project.query("SELECT 7").rows == [(7,)]


def test_timeout_interrupts_and_project_remains_usable(project):
    with pytest.raises(TimeoutError):
        project.query("SELECT sum(i) FROM range(1000000000) AS t(i)", timeout=0.02)
    assert project.query("SELECT 2").rows == [(2,)]


def test_user_cancel_interrupts_running_query(project):
    outcome = []

    def run():
        try:
            project.query("SELECT sum(i) FROM range(1000000000) AS t(i)", timeout=30)
        except Exception as exc:
            outcome.append(exc)

    worker = threading.Thread(target=run)
    worker.start()
    deadline = time.monotonic() + 2
    while project._active is None and time.monotonic() < deadline:
        time.sleep(0.001)
    assert project._active is not None
    project.cancel()
    worker.join(timeout=2)
    assert not worker.is_alive()
    assert len(outcome) == 1 and isinstance(outcome[0], InterruptedError)
    assert project.query("SELECT 4").rows == [(4,)]


def test_failed_export_preserves_existing_file(project, tmp_path):
    destination = tmp_path / "existing.csv"
    destination.write_text("keep me")
    with pytest.raises(FileExistsError):
        project.export("DELETE FROM _pip_sources", destination, "csv")
    assert destination.read_text() == "keep me"
    assert not list(tmp_path.glob(".pip-export-*"))


def test_export_cannot_overwrite_project_or_inputs(project, tmp_path):
    import os
    source = tmp_path / "original.csv"
    source.write_text("value\n7\n")
    project.import_file(source)
    alias = tmp_path / "hardlink.csv"
    os.link(source, alias)
    symlink = tmp_path / "symlink.csv"
    symlink.symlink_to(source)
    database_alias = tmp_path / "database-copy.parquet"
    os.link(project.path / "data.duckdb", database_alias)
    protected = [source, alias, symlink, database_alias, project.path / "data.duckdb", project.path / "project.json", project.path / "output.csv"]
    for path in protected:
        with pytest.raises(ValueError):
            project.export("SELECT 1 AS x", path, "csv")
    assert source.read_text() == "value\n7\n"
    assert project.query("SELECT value FROM original").rows == [(7,)]
    reopened = Project(project.path)
    assert reopened.tables()[0]["name"] == "original"
    reopened.close()


def test_export_never_replaces_any_existing_file(project, tmp_path):
    other = Project.create(tmp_path / "other")
    for target in [other.path / "project.json", other.path / "data.duckdb"]:
        with pytest.raises(FileExistsError):
            project.export("SELECT 1 AS x", target, "json")
    other.close()
    Project(tmp_path / "other").close()


def test_invalid_history_is_rejected_before_open(tmp_path):
    item = Project.create(tmp_path / "broken")
    item.close()
    meta = tmp_path / "broken" / "project.json"
    value = json.loads(meta.read_text())
    value["history"] = ["bad entry"]
    meta.write_text(json.dumps(value))
    with pytest.raises(ValueError, match="history"):
        Project(tmp_path / "broken")


def test_unreleased_older_project_format_is_rejected_unchanged(tmp_path):
    item = Project.create(tmp_path / "old")
    item.close()
    path = tmp_path / "old" / "project.json"
    meta = json.loads(path.read_text())
    meta["version"] = 1
    path.write_text(json.dumps(meta))
    original = path.read_bytes()
    with pytest.raises(ValueError, match="Unsupported"):
        Project(tmp_path / "old")
    assert path.read_bytes() == original


def test_structured_uuid_and_precise_values_export(project, tmp_path):
    sql = "SELECT '123e4567-e89b-12d3-a456-426614174000'::UUID AS id, [1,2] AS arr, 9007199254740993::BIGINT AS big, 0.10::DECIMAL(18,2) AS money"
    project.export(sql, tmp_path / "types.json", "json")
    data = json.loads((tmp_path / "types.json").read_text())[0]
    assert data["id"] == "123e4567-e89b-12d3-a456-426614174000"
    assert data["big"] == 9007199254740993 and data["money"] == "0.10"
    project.export(sql, tmp_path / "types.xlsx", "xlsx")
    book = load_workbook(tmp_path / "types.xlsx", read_only=True)
    values = list(book.active.values)[1]
    assert values == (data["id"], "[1, 2]", "9007199254740993", "0.10")
    book.close()


def test_mixed_large_integer_xlsx_keeps_exact_text(project, tmp_path):
    import zipfile
    path = tmp_path / "precision.xlsx"
    book = Workbook()
    book.active.append(["value"])
    book.active.append([9007199254740992])
    book.active.append([0.5])
    book.save(path)
    # Excel XML can contain integers beyond binary-double precision even
    # though openpyxl's writer itself rounds them when generating a fixture.
    with zipfile.ZipFile(path) as archive:
        entries = {entry.filename: archive.read(entry) for entry in archive.infolist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b"9007199254740992", b"9007199254740993")
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    table = project.import_file(path)
    assert project.query(f"SELECT value FROM {table}").rows == [("9007199254740993",), ("0.5",)]
