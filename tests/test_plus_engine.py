import datetime as dt
import threading
import time
from pathlib import Path

import pytest
from openpyxl import Workbook

from pip_tables import engine
from pip_tables.engine import Project


@pytest.fixture
def project(tmp_path):
    item = Project.create(tmp_path / "project")
    yield item
    item.close()


def test_inspection_matches_import_and_does_not_change_project(project, tmp_path, monkeypatch):
    source = tmp_path / "input.csv"
    source.write_text("key,amount\n" + "".join(f"{i},{i + 0.25}\n" for i in range(25)), encoding="utf-8")
    before = {name: (project.path / name).read_bytes() for name in ("data.duckdb", "project.json")}
    folders = []
    original = engine.tempfile.TemporaryDirectory

    def tracked(*args, **kwargs):
        kwargs["dir"] = tmp_path
        item = original(*args, **kwargs)
        folders.append(Path(item.name))
        return item

    monkeypatch.setattr(engine.tempfile, "TemporaryDirectory", tracked)
    inspected = project.inspect_file(source)
    assert inspected["rows"] == [(i, i + 0.25) for i in range(20)]
    assert inspected["total"] == 25 and inspected["sheets"] == []
    assert inspected["fingerprint"].startswith("sha256:")
    assert all(not folder.exists() for folder in folders)
    assert {name: (project.path / name).read_bytes() for name in before} == before
    assert project.tables() == [] and project.history() == []
    name = project.import_file(source, expected_fingerprint=inspected["fingerprint"])
    assert project.tables()[0]["columns"] == inspected["columns"]
    assert project.query(f"SELECT * FROM {name} ORDER BY key", limit=20).rows == inspected["rows"]


def test_sheet_names_selection_and_exact_precision(project, tmp_path):
    source = tmp_path / "ledger.xlsx"
    book = Workbook()
    first = book.active
    first.title = "Summary"
    first.append(["id", "value"])
    first.append([1, 3])
    detail = book.create_sheet("Detail")
    detail.append(["date", "wide", "money"])
    detail.append([dt.date(2026, 10, 23), 9007199254740993, "0.10"])
    book.save(source)
    assert project.worksheet_names(source) == ["Summary", "Detail"]
    chooser = project.inspect_file(source)
    assert chooser["sheets"] == ["Summary", "Detail"]
    assert chooser["columns"] == [] and chooser["rows"] == [] and chooser["total"] == 0
    selected = project.inspect_file(source, sheet="Detail")
    assert selected["columns"][1]["type"] in {"BIGINT", "VARCHAR"}
    name = project.import_file(source, sheet="Detail", expected_fingerprint=selected["fingerprint"])
    assert project.tables()[0]["columns"] == selected["columns"]
    assert project.query(f"SELECT * FROM {name}").rows == selected["rows"]
    assert selected["rows"][0][2] == "0.10"


def test_sheet_picker_reads_only_bounded_workbook_metadata(project, tmp_path, monkeypatch):
    source = tmp_path / "sheets.xlsx"
    book = Workbook()
    book.active.title = "One"
    book.create_sheet("Two")
    book.save(source)

    def must_not_open(*_args, **_kwargs):
        raise AssertionError("Sheet picker must not parse workbook cells or shared strings")

    monkeypatch.setattr(engine, "load_workbook", must_not_open)
    assert project.worksheet_names(source) == ["One", "Two"]
    assert project.inspect_file(source)["sheets"] == ["One", "Two"]


def test_inspection_rejects_formula_without_changing_project(project, tmp_path):
    source = tmp_path / "formula.xlsx"
    book = Workbook()
    book.active.append(["value"])
    book.active.append(["=1+1"])
    book.save(source)
    with pytest.raises(ValueError, match="formulas"):
        project.inspect_file(source)
    assert project.tables() == []


def test_fingerprint_blocks_modified_source_before_and_during_import(project, tmp_path, monkeypatch):
    source = tmp_path / "data.csv"
    source.write_text("value\n1\n")
    stamp = project.inspect_file(source)["fingerprint"]
    source.write_text("value\n2\n")
    with pytest.raises(ValueError, match="changed"):
        project.import_file(source, expected_fingerprint=stamp)
    assert project.tables() == []

    stamp = project.inspect_file(source)["fingerprint"]
    loader = Project._load_file

    def changed(conn, path, kind, name, sheet, encoding, all_text, cancelled):
        loader(conn, path, kind, name, sheet, encoding, all_text, cancelled)
        path.write_text("value\n3\n")

    monkeypatch.setattr(Project, "_load_file", staticmethod(changed))
    with pytest.raises(ValueError, match="changed during import"):
        project.import_file(source, expected_fingerprint=stamp)
    assert project.tables() == []


def test_inspection_cancel_cleans_up(project, tmp_path, monkeypatch):
    source = tmp_path / "large.csv"
    source.write_text("value\n1\n")
    entered = threading.Event()
    release = threading.Event()
    original = Project._fingerprint

    def held(self, path):
        entered.set()
        release.wait(2)
        return original(self, path)

    monkeypatch.setattr(Project, "_fingerprint", held)
    outcome = []

    def inspect():
        try:
            project.inspect_file(source)
        except Exception as exc:
            outcome.append(exc)

    worker = threading.Thread(target=inspect)
    worker.start()
    assert entered.wait(2)
    project.cancel()
    release.set()
    worker.join(2)
    assert not worker.is_alive()
    assert len(outcome) == 1 and isinstance(outcome[0], InterruptedError)
    assert project.tables() == []
    assert project.inspect_file(source)["total"] == 1


def test_pages_have_full_count_offset_bounds_and_no_history_flood(project):
    sql = "SELECT i FROM range(23) AS t(i) ORDER BY i"
    first = project.query(sql, limit=10)
    second = project.query(sql, limit=10, offset=10, record_history=False)
    last = project.query(sql, limit=10, offset=20, record_history=False)
    beyond = project.query(sql, limit=10, offset=30, record_history=False)
    assert [r[0] for r in first.rows + second.rows + last.rows] == list(range(23))
    assert [r.total for r in (first, second, last, beyond)] == [23] * 4
    assert [r.offset for r in (first, second, last, beyond)] == [0, 10, 20, 30]
    assert [r.truncated for r in (first, second, last, beyond)] == [True, True, False, False]
    assert beyond.rows == [] and len(project.history()) == 1
    for bad in (-1, True, 1.2, 2**63):
        with pytest.raises(ValueError, match="Offset"):
            project.query(sql, offset=bad)
